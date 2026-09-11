"""Report exports: a full SQL dump, and an Excel workbook of the panel's two views.

openpyxl, not pandas: install.sh skips pandas on Termux (Android), and this
has to work there too.
"""
import io

from openpyxl import Workbook
from openpyxl.styles import Font

import db
from confidence import STATUS_ES

_OBS_COLUMNS = [
    ("id", "ID"), ("customer", "Cliente"), ("city", "Ciudad"), ("country", "País"),
    ("modality", "Modalidad"), ("brand", "Marca"), ("model", "Modelo"),
    ("quantity", "Cantidad"), ("age_years", "Antigüedad (años)"), ("status", "Estado"),
    ("confidence", "Confianza"), ("corroboration_count", "Corroboraciones"),
    ("created_at", "Fecha (UTC)"), ("source_text", "Texto fuente"),
]

# Same grouping as the panel's "Vista agregada entre clientes".
_AGGREGATE_SQL = """
    SELECT modality, brand, COALESCE(SUM(quantity), 0), ROUND(AVG(age_years), 1),
           COUNT(DISTINCT customer)
    FROM observations GROUP BY modality, brand ORDER BY 3 DESC
"""
_AGGREGATE_HEADERS = ["Modalidad", "Marca", "Cantidad total", "Antigüedad promedio", "N.º clientes"]


def sql_dump(conn) -> str:
    # Whole DB (observations + taxonomy), restorable with `sqlite3 new.db < file.sql`.
    return "\n".join(conn.iterdump()) + "\n"


def _fill_sheet(ws, headers: list[str], rows) -> None:
    ws.append(headers)
    for row in rows:
        ws.append(list(row))
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"
    for col in ws.columns:
        for cell in col:
            # Names/narration are user input: a leading "=" would otherwise be
            # written as a live formula (spreadsheet formula injection).
            if isinstance(cell.value, str) and cell.value.startswith("="):
                cell.data_type = "s"
        width = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[col[0].column_letter].width = min(width + 2, 60)


def xlsx(conn) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Observaciones"
    _fill_sheet(ws, [label for _, label in _OBS_COLUMNS], (
        [STATUS_ES.get(o[k], o[k]) if k == "status" else o[k] for k, _ in _OBS_COLUMNS]
        for o in db.get_all_observations(conn)
    ))
    _fill_sheet(wb.create_sheet("Agregado"), _AGGREGATE_HEADERS, conn.execute(_AGGREGATE_SQL))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

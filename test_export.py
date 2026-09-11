import io
import os
import sqlite3
import tempfile
import unittest

from openpyxl import load_workbook

import db
import export


class TestExport(unittest.TestCase):
    def setUp(self):
        fd, self.path = tempfile.mkstemp(suffix=".db")
        os.close(fd)
        os.remove(self.path)  # must not exist yet, so init_db treats it as new
        self.conn = db.init_db(self.path)
        for customer, qty in [("Hospital A", 2), ("=HYPERLINK(\"http://x\")", 1)]:
            db.insert_observation(self.conn, {
                "customer": customer, "modality": "Monitor de Paciente", "brand": "Philips",
                "quantity": qty, "status": "reported", "confidence": 0.7,
            })

    def tearDown(self):
        self.conn.close()
        os.remove(self.path)

    def test_xlsx_has_both_views(self):
        wb = load_workbook(io.BytesIO(export.xlsx(self.conn)))
        self.assertEqual(wb.sheetnames, ["Observaciones", "Agregado"])
        obs = wb["Observaciones"]
        self.assertEqual(obs.max_row, 3)  # header + 2 observations
        self.assertIn("Reportado", [c.value for c in obs[2]])
        agg = wb["Agregado"]
        self.assertEqual([c.value for c in agg[2]], ["Monitor de Paciente", "Philips", 3, None, 2])

    def test_xlsx_never_writes_formulas(self):
        wb = load_workbook(io.BytesIO(export.xlsx(self.conn)))
        cells = [c for row in wb["Observaciones"].iter_rows() for c in row]
        self.assertFalse([c.coordinate for c in cells if c.data_type == "f"])

    def test_sql_dump_restores(self):
        restored = sqlite3.connect(":memory:")
        restored.executescript(export.sql_dump(self.conn))
        self.assertEqual(restored.execute("SELECT COUNT(*) FROM observations").fetchone()[0], 2)
        self.assertTrue(restored.execute("SELECT COUNT(*) FROM taxonomy").fetchone()[0] > 0)


if __name__ == "__main__":
    unittest.main()
